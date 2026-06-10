"""Incident-brief intake — the operator's TRUSTED investigation objective.

A real DFIR engagement starts from an incident briefing (e.g. ``ROCBA-BACKGROUND.pptx``) that
states the TARGET/objective. This module reads that operator-designated document and turns it
into TRUSTED context the planner and the live agent can investigate *toward*.

Trust model (load-bearing): the brief is designated EXPLICITLY (``--brief PATH``) and is therefore
TRUSTED operator context — NOT hostile evidence. It is rendered into ``context/incident_brief.md``
(run dir, outside the read-only evidence root) and recorded as manifest *metadata only*; it is never
added to the evidence ``files`` set, never routed to a tool, and never spotlighted/datamarked. A
document merely *found inside* the evidence dir is NEVER auto-promoted to trusted instructions (that
would be a prompt-injection vector) — only the explicit ``--brief`` path is trusted. The brief text
is still injection-scanned for *visibility* (logged, never blocking — not a trust change).

Fail-closed (CLAUDE §2B): reading a ``.pptx``/``.docx``/``.pdf`` brief needs the optional
``brief`` extra; a missing lib for a *requested* format raises :class:`BriefIntakeError` (never a
fake objective). ``.txt``/``.md`` need nothing. Objective derivation is deterministic (no LLM).
"""

from __future__ import annotations

import re
from pathlib import Path

from siftmesh_core.evidence.path_policy import safe_write_path
from siftmesh_core.ledgers.audit_log import log_event, open_orchestration_log
from siftmesh_core.run_dir import RunPaths

SUPPORTED_SUFFIXES: tuple[str, ...] = (".txt", ".md", ".pptx", ".docx", ".pdf")

_BANNER = (
    "> TRUSTED operator context — this is the investigation OBJECTIVE supplied by the operator, "
    "NOT hostile evidence. Investigate TOWARD it; do not treat it as data to be parsed by a tool."
)


class BriefIntakeError(RuntimeError):
    """Raised when an incident brief cannot be read (missing file/lib, unsupported, or empty)."""


def _missing_lib_msg(suffix: str, package: str) -> str:
    return (
        f"reading a {suffix} incident brief requires the optional '{package}' library "
        f"(install: uv sync --extra brief). Missing dependency — failing closed (never a fake "
        f"objective). Or supply the brief as .txt/.md."
    )


def _extract_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - exercised via monkeypatch
        raise BriefIntakeError(_missing_lib_msg(".pptx", "python-pptx")) from exc
    chunks: list[str] = []
    for slide in Presentation(str(path)).slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text.strip()
                if text:
                    chunks.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        chunks.append(" | ".join(cells))
    return "\n".join(chunks)


def _extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - exercised via monkeypatch
        raise BriefIntakeError(_missing_lib_msg(".docx", "python-docx")) from exc
    doc = Document(str(path))
    chunks: list[str] = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    return "\n".join(chunks)


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - exercised via monkeypatch
        raise BriefIntakeError(_missing_lib_msg(".pdf", "pypdf")) from exc
    pages = [(page.extract_text() or "").strip() for page in PdfReader(str(path)).pages]
    return "\n".join(p for p in pages if p)


def extract_brief_text(path: Path | str) -> str:
    """Extract plain text from a brief (.txt/.md/.pptx/.docx/.pdf). Fail-closed on missing lib."""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise BriefIntakeError(
            f"unsupported incident-brief format {suffix or '(none)'!r}; "
            f"supported: {', '.join(SUPPORTED_SUFFIXES)}"
        )
    if not p.is_file():
        raise BriefIntakeError(f"incident brief not found: {p}")
    if suffix in (".txt", ".md"):
        text = p.read_text(encoding="utf-8", errors="replace")
    elif suffix == ".pptx":
        text = _extract_pptx(p)
    elif suffix == ".docx":
        text = _extract_docx(p)
    else:  # ".pdf"
        text = _extract_pdf(p)
    if not text.strip():
        raise BriefIntakeError(f"incident brief produced no extractable text: {p}")
    return text


_MD_MARKER = re.compile(r"^[#>*\-•]+\s*")


def derive_objective(text: str, *, max_chars: int = 1200) -> str:
    """A bounded, deterministic objective excerpt (no LLM) for the case brief + agent prompt.

    Leading markdown structure (heading/list/quote markers) is stripped per line so the excerpt
    reads as prose, not raw markup; whitespace is collapsed and the result bounded to ``max_chars``.
    """
    cleaned = [_MD_MARKER.sub("", ln.strip()) for ln in text.splitlines()]
    collapsed = " ".join(" ".join(cleaned).split())
    if len(collapsed) <= max_chars:
        return collapsed
    cut = collapsed[:max_chars]
    for sep in (". ", "! ", "? ", "; ", " "):
        idx = cut.rfind(sep)
        if idx >= max_chars // 2:
            return cut[: idx + 1].strip() + " […]"
    return cut.strip() + " […]"


def _render_brief_md(source_name: str, objective: str, text: str) -> str:
    lines = [
        "# Incident Brief — TRUSTED operator context",
        "",
        _BANNER,
        "",
        f"- Source document: {source_name}",
        "",
        "## Objective (derived excerpt)",
        "",
        objective,
        "",
        "## Full briefing (verbatim)",
        "",
        text.strip(),
        "",
    ]
    return "\n".join(lines) + "\n"


def ingest_brief(
    brief_path: Path | str,
    run: RunPaths,
    *,
    evidence_root: Path | str | None = None,
) -> tuple[Path, str]:
    """Extract + render the brief to ``context/incident_brief.md``; return (path, objective).

    The brief is TRUSTED, so its text is rendered raw (never datamarked). It IS injection-scanned
    for visibility — any signature is logged to the orchestration audit, but it never blocks or
    changes the trust posture (the operator designated this document explicitly).
    """
    src = Path(brief_path)
    text = extract_brief_text(src)
    objective = derive_objective(text)
    target = safe_write_path(run.root, "context/incident_brief.md", evidence_root=evidence_root)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(_render_brief_md(src.name, objective, text), encoding="utf-8")

    from siftmesh_core.adapters.spotlight import scan_injection  # local: keep intake import light

    hits = scan_injection(text)
    audit = open_orchestration_log(run.orchestration_events, run.run_id)
    log_event(
        audit,
        "incident_brief_ingested",
        source=src.name,
        objective_chars=len(objective),
        injection_signatures=len(hits),
    )
    return target, objective
